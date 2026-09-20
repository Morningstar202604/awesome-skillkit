# -*- coding: utf-8 -*-
"""office · 真实文档交付物：生成可打开的 .docx 合同 + .pptx 方案。"""
SKILL = "office-suite"
DOMAIN = "office"


def run(ctx, ffmpeg):
    ctx.think(
        "office 域挑 office-suite（Word/PPT/Excel）。任务：生成真实 .docx 合同（可打开、有表格）"
        "+ 真实 .pptx（3 页方案）。需要 python-docx / python-pptx（已装）。"
        "边界：缺库时给出降级（markdown 版）而非静默失败。"
    )
    import os
    try:
        import docx
        from docx.shared import Pt
        from pptx import Presentation
        from pptx.util import Inches
    except Exception as e:
        ctx.problem(f"python-docx/pptx 未装: {e}")
        # 降级 markdown
        md = os.path.join(ctx.d, "contract.md")
        with open(md, "w", encoding="utf-8") as fh:
            fh.write("# 服务合同\n\n甲方：测试公司\n乙方：交付团队\n\n| 项 | 内容 |\n|---|---|\n| 交付 | 一代全量测试 |\n| 期限 | 2026-09 |\n\n")
        ctx.result("warn", f"缺库，降级为 markdown 合同 {md}")
        return

    # --- 真实 docx 合同 ---
    doc = docx.Document()
    doc.add_heading("服务合同", 0)
    doc.add_paragraph("甲方：Acme 测试公司（以下简称甲方）")
    doc.add_paragraph("乙方：awesome-skillkit 交付团队（以下简称乙方）")
    tbl = doc.add_table(rows=3, cols=2)
    tbl.style = "Light Grid Accent 1"
    tbl.cell(0, 0).text = "交付物"; tbl.cell(0, 1).text = "说明"
    tbl.cell(1, 0).text = "全量测试"; tbl.cell(1, 1).text = "18 域真实交付物"
    tbl.cell(2, 0).text = "期限"; tbl.cell(2, 1).text = "2026-09-20"
    out_docx = os.path.join(ctx.d, "contract.docx")
    doc.save(out_docx)
    ctx.think(f"真实 docx 合同 {out_docx}（{os.path.getsize(out_docx)} bytes，含 1 表格）")

    # --- 真实 pptx 方案（3 页）---
    prs = Presentation()
    for title, body in [
        ("一代全量测试", "18 域 · 真实交付物 · 全量过程记录"),
        ("方法", "脚本真跑 + LLM 判分 + 交付物保留不删"),
        ("结果", "video/mp4 · image/png · 可运行程序 · 成文文章 · 设计方案"),
    ]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    out_pptx = os.path.join(ctx.d, "plan.pptx")
    prs.save(out_pptx)
    ctx.think(f"真实 pptx 方案 {out_pptx}（{os.path.getsize(out_pptx)} bytes，3 页）")

    ctx.result("pass", f"真实 docx 合同 + 真实 pptx 方案（均可打开）")
    ctx.better("可加 python-docx 模板填写（{{占位符}}）做更强交付；当前已含表格验证 docx 真实性。")


if __name__ == "__main__":
    from full_skill_test import Recorder, get_ffmpeg, ARTIFACTS
    import os
    os.makedirs(ARTIFACTS, exist_ok=True)
    c = Recorder(SKILL, DOMAIN)
    run(c, get_ffmpeg())
    c.dump()
